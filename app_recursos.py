import streamlit as st
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_ORIENTATION
from docx.oxml import OxmlElement
from datetime import date
import locale
import os

# Configurar el locale a español para las fechas
locale.setlocale(locale.LC_TIME, 'es_ES.UTF-8')

# Configurar la página de Streamlit
st.set_page_config(page_title="Generador de Planificaciones", page_icon="📚")

st.title("📚 Generador de Planificaciones Diarias para Docentes")

# Inicializar session_state para almacenar planificaciones
if "planificaciones" not in st.session_state:
    st.session_state.planificaciones = []

# Campos para el encabezado
semana = st.text_input("Semana", "")
curso = st.selectbox("Curso", ["1ro de Primaria", "2do de Primaria", "3ro de Primaria", 
                               "4to de Primaria", "5to de Primaria", "6to de Primaria"])
materia = st.text_input("Materia", "")
profesor = st.text_input("Profesor/a", "")

# Competencias
competencias_fundamentales = st.text_area("Competencias Fundamentales", "")

# Seleccionar fecha con un calendario
fecha = st.date_input("Fecha", date.today())
fecha_formateada = fecha.strftime("%A, %d de %B").capitalize()

# Datos para un día
st.subheader("📅 Planificación Diaria")
competencias_especificas = st.text_area("Competencias Específicas", "")
contenido = st.text_area("Contenido", "")
indicadores_logros = st.text_area("Indicadores de Logros", "")
secuencia_actividades = st.text_area("Secuencia de Actividades", "")
ejes_tematicos = st.text_input("Ejes Temáticos Transversales", "")
tipo_evaluacion = st.text_input("Tipo de Evaluación", "")
asignaciones = st.text_area("Asignaciones", "")

# Campo para subir imágenes
st.subheader("📂 Recursos")
imagenes = st.file_uploader("Subir imágenes (opcional)", type=["png", "jpg", "jpeg"], accept_multiple_files=True)

# Campo para ingresar links externos
links = st.text_area("Enlaces externos (opcional, uno por línea)")

# Botón para agregar la planificación del día a la lista
if st.button("Agregar Planificación del Día"):
    recursos = []
    if imagenes:
        recursos.extend([f"Imagen: {img.name}" for img in imagenes])
    if links:
        recursos.extend([f"Link: {link}" for link in links.splitlines()])
    
    planificacion_dia = {
        "fecha": fecha_formateada,
        "competencias_especificas": competencias_especificas,
        "contenido": contenido,
        "indicadores_logros": indicadores_logros,
        "secuencia_actividades": secuencia_actividades,
        "ejes_tematicos": ejes_tematicos,
        "tipo_evaluacion": tipo_evaluacion,
        "asignaciones": asignaciones,
        "recursos": recursos,
        "imagenes": imagenes,  # Guardar imágenes
        "links": links.splitlines()  # Guardar links como lista
    }
    st.session_state.planificaciones.append(planificacion_dia)
    st.success(f"¡Planificación agregada para el día {fecha_formateada}!")

# Mostrar planificaciones registradas
st.subheader("📅 Planificaciones Registradas")
if st.session_state.planificaciones:
    for idx, plan in enumerate(st.session_state.planificaciones, 1):
        st.write(f"**Día {idx}: {plan['fecha']}**")
        st.write(f"- **Competencias Específicas:** {plan['competencias_especificas']}")
        st.write(f"- **Contenido:** {plan['contenido']}")
        st.write(f"- **Indicadores de Logros:** {plan['indicadores_logros']}")
        st.write(f"- **Secuencia de Actividades:** {plan['secuencia_actividades']}")
        st.write(f"- **Ejes Temáticos:** {plan['ejes_tematicos']}")
        st.write(f"- **Tipo de Evaluación:** {plan['tipo_evaluacion']}")
        st.write(f"- **Asignaciones:** {plan['asignaciones']}")
        if plan['recursos']:
            st.write(f"- **Recursos:**")
            for recurso in plan['recursos']:
                st.write(f"  - {recurso}")
        st.write("---")
else:
    st.info("No hay planificaciones registradas aún.")

# Botón para generar el documento consolidado
if st.button("Generar Documento Consolidado"):
    if st.session_state.planificaciones:
        doc = Document()
        doc.add_heading('Planificación Semanal Consolidada', level=1)
        doc.add_paragraph(f"Semana: {semana}        Curso: {curso}        Materia: {materia}        Profesor/a: {profesor}")
        
        doc.add_paragraph("Competencias Fundamentales:").bold = True
        doc.add_paragraph(competencias_fundamentales)
        
        # Cambiar la orientación a horizontal
        seccion = doc.sections[0]
        seccion.orientation = WD_ORIENTATION.LANDSCAPE
        seccion.page_width, seccion.page_height = seccion.page_height, seccion.page_width
        
        # Crear tabla con columna "Recursos"
        tabla = doc.add_table(rows=1, cols=9)
        tabla.style = 'Table Grid'
        encabezados = ["Fecha", "Competencias Específicas", "Contenidos", "Indicadores de Logros", 
                       "Secuencia de Actividades", "Ejes Temáticos Transversales", "Tipo de Evaluación", 
                       "Asignaciones", "Recursos"]
        
        # Encabezados con estilo
        hdr_cells = tabla.rows[0].cells
        for i, encabezado in enumerate(encabezados):
            hdr_cells[i].text = encabezado
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(12)
        
        # Agregar filas con la información de cada día
        for plan in st.session_state.planificaciones:
            fila = tabla.add_row().cells
            fila[0].text = plan["fecha"]
            fila[1].text = plan["competencias_especificas"]
            fila[2].text = plan["contenido"]
            fila[3].text = plan["indicadores_logros"]
            fila[4].text = plan["secuencia_actividades"]
            fila[5].text = plan["ejes_tematicos"]
            fila[6].text = plan["tipo_evaluacion"]
            fila[7].text = plan["asignaciones"]

            # Agregar recursos a la columna de "Recursos"
            recursos_texto = ""
            if plan["recursos"]:
                for recurso in plan["recursos"]:
                    recursos_texto += f"{recurso}\n"
            else:
                recursos_texto = "N/A"
            fila[8].text = recursos_texto.strip()
        
        # Guardar el documento
        doc.save('planificacion_consolidada.docx')
        st.success("¡Documento consolidado generado con éxito!")
        with open('planificacion_consolidada.docx', 'rb') as f:
            st.download_button("Descargar Planificación Consolidada", f, file_name="planificacion_consolidada.docx")
    else:
        st.warning("No hay planificaciones registradas para generar el documento.")
